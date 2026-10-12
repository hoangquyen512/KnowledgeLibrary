from __future__ import annotations

import io
from typing import Any

from pptx import Presentation
from pptx.util import Inches, Pt


def export_presentation_pptx(payload: dict[str, Any], *, deck_title: str) -> bytes:
    prs = Presentation()
    prs.core_properties.title = deck_title[:255]

    slides = payload.get("slides") or []
    if not slides:
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = deck_title
        return _save(prs)

    for idx, row in enumerate(slides):
        layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]
        slide = prs.slides.add_slide(layout)
        title_shape = slide.shapes.title
        if title_shape:
            title_shape.text = str(row.get("title") or f"Slide {idx + 1}")

        body = slide.placeholders[1] if len(slide.placeholders) > 1 else None
        bullets = row.get("bullets") or []
        if body and bullets:
            tf = body.text_frame
            tf.clear()
            for i, bullet in enumerate(bullets):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = str(bullet)
                p.level = 0
                p.font.size = Pt(18)

        notes = str(row.get("notes") or "").strip()
        if notes:
            notes_slide = slide.notes_slide
            notes_slide.notes_text_frame.text = notes

    return _save(prs)


def _save(prs: Presentation) -> bytes:
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
