"""Excel (.xlsx / .xls) parser → sheet documents with row-range chunks."""

from __future__ import annotations

import io
from typing import Any

from app.ingestion.parsers import (
    ChunkHint,
    NormalizedDocument,
    NormalizedSection,
    ParseResult,
    register_parser,
)
from app.models.document import DocumentType
from app.models.source import SourceType

_ROWS_PER_CHUNK = 40


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _sheet_to_rows(ws: Any) -> tuple[list[str], list[list[str]]]:
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_raw = next(rows_iter)
    except StopIteration:
        return [], []
    headers = [_cell_str(c) or f"Column{i+1}" for i, c in enumerate(header_raw)]
    data: list[list[str]] = []
    for row in rows_iter:
        values = [_cell_str(c) for c in row]
        if not any(values):
            continue
        # pad / trim to header length
        if len(values) < len(headers):
            values.extend([""] * (len(headers) - len(values)))
        data.append(values[: len(headers)])
    return headers, data


def _format_table(headers: list[str], rows: list[list[str]], *, start_row: int) -> str:
    lines = [" | ".join(headers), " | ".join("---" for _ in headers)]
    for _offset, row in enumerate(rows):
        lines.append(" | ".join(row))
    return f"Rows {start_row}–{start_row + len(rows) - 1}\n" + "\n".join(lines)


class ExcelParser:
    source_type = SourceType.FILE_EXCEL

    def parse(
        self, data: bytes, *, filename: str, options: dict[str, Any] | None = None
    ) -> ParseResult:
        options = options or {}
        selected_sheets: set[str] | None = None
        if options.get("sheet_names"):
            selected_sheets = {str(s) for s in options["sheet_names"]}

        lower = filename.lower()
        try:
            if lower.endswith(".xls") and not lower.endswith(".xlsx"):
                return self._parse_xls(data, filename=filename, selected_sheets=selected_sheets)
            return self._parse_xlsx(data, filename=filename, selected_sheets=selected_sheets)
        except Exception as exc:  # noqa: BLE001
            return ParseResult(
                documents=[],
                error_code="EXCEL_PARSE_FAILED",
                error_message=str(exc)[:500],
            )

    def _parse_xlsx(
        self,
        data: bytes,
        *,
        filename: str,
        selected_sheets: set[str] | None,
    ) -> ParseResult:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet_names = list(wb.sheetnames)
        preview_sheets: list[dict[str, Any]] = []
        documents: list[NormalizedDocument] = []

        for name in sheet_names:
            ws = wb[name]
            headers, rows = _sheet_to_rows(ws)
            empty = len(rows) == 0 and not any(headers)
            preview_sheets.append(
                {
                    "name": name,
                    "row_count": len(rows),
                    "column_count": len(headers),
                    "empty": empty,
                    "selected": (selected_sheets is None and not empty)
                    or (selected_sheets is not None and name in selected_sheets),
                }
            )
            if empty:
                continue
            if selected_sheets is not None and name not in selected_sheets:
                continue
            documents.append(self._build_doc(filename, name, headers, rows))

        wb.close()
        return ParseResult(
            documents=documents,
            preview={
                "filename": filename,
                "sheets": preview_sheets,
                "workbook_type": "xlsx",
            },
        )

    def _parse_xls(
        self,
        data: bytes,
        *,
        filename: str,
        selected_sheets: set[str] | None,
    ) -> ParseResult:
        try:
            import xlrd
        except ImportError:
            return ParseResult(
                documents=[],
                error_code="EXCEL_PARSE_FAILED",
                error_message="xlrd is required to parse .xls files.",
            )
        book = xlrd.open_workbook(file_contents=data)
        preview_sheets: list[dict[str, Any]] = []
        documents: list[NormalizedDocument] = []
        for idx in range(book.nsheets):
            sheet = book.sheet_by_index(idx)
            name = sheet.name
            headers: list[str] = []
            rows: list[list[str]] = []
            if sheet.nrows > 0:
                headers = [
                    _cell_str(sheet.cell_value(0, c)) or f"Column{c+1}"
                    for c in range(sheet.ncols)
                ]
                for r in range(1, sheet.nrows):
                    values = [_cell_str(sheet.cell_value(r, c)) for c in range(sheet.ncols)]
                    if any(values):
                        rows.append(values)
            empty = len(rows) == 0 and not any(headers)
            preview_sheets.append(
                {
                    "name": name,
                    "row_count": len(rows),
                    "column_count": len(headers),
                    "empty": empty,
                    "selected": (selected_sheets is None and not empty)
                    or (selected_sheets is not None and name in selected_sheets),
                }
            )
            if empty:
                continue
            if selected_sheets is not None and name not in selected_sheets:
                continue
            documents.append(self._build_doc(filename, name, headers, rows))
        return ParseResult(
            documents=documents,
            preview={"filename": filename, "sheets": preview_sheets, "workbook_type": "xls"},
        )

    def _build_doc(
        self,
        filename: str,
        sheet_name: str,
        headers: list[str],
        rows: list[list[str]],
    ) -> NormalizedDocument:
        header_line = "Columns: " + " | ".join(headers)
        body_parts = [f"Sheet: {sheet_name}", header_line, ""]
        for i, row in enumerate(rows, start=1):
            body_parts.append(f"Row {i}: " + " | ".join(row))
        text = "\n".join(body_parts)

        chunks: list[ChunkHint] = []
        for start in range(0, max(len(rows), 1), _ROWS_PER_CHUNK):
            batch = rows[start : start + _ROWS_PER_CHUNK]
            if not batch and start > 0:
                break
            row_start = start + 1
            row_end = start + len(batch)
            chunk_text = (
                f"Sheet: {sheet_name}\n{header_line}\n\n"
                + _format_table(headers, batch, start_row=row_start)
            )
            chunks.append(
                ChunkHint(
                    text=chunk_text,
                    position=len(chunks),
                    metadata={
                        "sheet": sheet_name,
                        "row_start": row_start,
                        "row_end": row_end,
                        "citation_label": f"{filename} — Sheet {sheet_name} — Rows {row_start}–{row_end}",
                    },
                )
            )

        return NormalizedDocument(
            title=f"Sheet: {sheet_name}",
            document_type=DocumentType.SHEET,
            source_type=SourceType.FILE_EXCEL,
            source_locator=f"{filename}#{sheet_name}",
            normalized_text=text,
            external_id=f"sheet:{sheet_name}",
            sections=[NormalizedSection(title=sheet_name, text=text)],
            chunk_hints=chunks,
            raw_metadata={
                "sheet_name": sheet_name,
                "row_count": len(rows),
                "column_count": len(headers),
                "filename": filename,
            },
        )


register_parser(ExcelParser())
