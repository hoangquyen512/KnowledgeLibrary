"""PowerPoint (.pptx) parser — one document, slide chunks."""

from __future__ import annotations

from typing import Any

from app.ingestion.parsers import (
    ChunkHint,
    NormalizedDocument,
    NormalizedSection,
    ParseResult,
    register_parser,
)
from app.models.document import DocumentType
from app.models.source import SourceType


class PptxParser:
    source_type = SourceType.FILE_POWERPOINT

    def parse(
        self, data: bytes, *, filename: str, options: dict[str, Any] | None = None
    ) -> ParseResult:
        try:
            import io

            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE_TYPE

            prs = Presentation(io.BytesIO(data))
            sections: list[NormalizedSection] = []
            chunks: list[ChunkHint] = []
            all_parts: list[str] = []
            image_count = 0

            for idx, slide in enumerate(prs.slides, start=1):
                texts: list[str] = []
                title = ""
                for shape in slide.shapes:
                    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                        image_count += 1
                    if not shape.has_text_frame:
                        if shape.has_table:
                            table_rows: list[str] = []
                            for row in shape.table.rows:
                                cells = [c.text.strip() for c in row.cells]
                                table_rows.append(" | ".join(cells))
                            if table_rows:
                                texts.append("Table:\n" + "\n".join(table_rows))
                        continue
                    shape_text = (shape.text or "").strip()
                    if not shape_text:
                        continue
                    if not title and (
                        getattr(shape, "is_placeholder", False)
                        or "title" in (shape.name or "").lower()
                    ):
                        title = shape_text.split("\n", 1)[0][:200]
                    texts.append(shape_text)

                notes = ""
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes = (slide.notes_slide.notes_text_frame.text or "").strip()

                slide_title = title or f"Slide {idx}"
                body = "\n".join(texts).strip()
                parts = [f"Slide {idx}", f"Title: {slide_title}", "", "Content:", body or "(empty)"]
                if notes:
                    parts.extend(["", "Speaker Notes:", notes])
                slide_text = "\n".join(parts)
                all_parts.append(slide_text)
                sections.append(NormalizedSection(title=slide_title, text=slide_text))
                chunks.append(
                    ChunkHint(
                        text=slide_text,
                        position=len(chunks),
                        metadata={
                            "slide_number": idx,
                            "slide_title": slide_title,
                            "citation_label": f"{filename} — Slide {idx}",
                        },
                    )
                )

            # Merge tiny consecutive slides
            merged: list[ChunkHint] = []
            buf: list[ChunkHint] = []

            def flush_buf() -> None:
                nonlocal buf
                if not buf:
                    return
                if len(buf) == 1:
                    merged.append(buf[0])
                else:
                    nums = [c.metadata.get("slide_number") for c in buf]
                    text = "\n\n".join(c.text for c in buf)
                    first = buf[0]
                    last = buf[-1]
                    merged.append(
                        ChunkHint(
                            text=text,
                            position=len(merged),
                            metadata={
                                "slide_number": first.metadata.get("slide_number"),
                                "slide_numbers": nums,
                                "slide_title": first.metadata.get("slide_title"),
                                "citation_label": (
                                    f"{filename} — Slides {nums[0]}–{nums[-1]}"
                                    if nums[0] != nums[-1]
                                    else first.metadata.get("citation_label")
                                ),
                            },
                        )
                    )
                buf = []

            for chunk in chunks:
                if len(chunk.text) < 120:
                    buf.append(chunk)
                    if len(buf) >= 3:
                        flush_buf()
                else:
                    flush_buf()
                    merged.append(
                        ChunkHint(text=chunk.text, position=len(merged), metadata=chunk.metadata)
                    )
            flush_buf()

            title = filename.rsplit(".", 1)[0] or filename
            return ParseResult(
                documents=[
                    NormalizedDocument(
                        title=title,
                        document_type=DocumentType.PRESENTATION,
                        source_type=SourceType.FILE_POWERPOINT,
                        source_locator=filename,
                        normalized_text="\n\n".join(all_parts),
                        external_id=f"pptx:{filename}",
                        sections=sections,
                        chunk_hints=merged or chunks,
                        raw_metadata={
                            "filename": filename,
                            "slide_count": len(sections),
                            "image_count": image_count,
                        },
                    )
                ],
                preview={
                    "filename": filename,
                    "title": title,
                    "slide_count": len(sections),
                    "image_count": image_count,
                },
            )
        except Exception as exc:  # noqa: BLE001
            return ParseResult(
                documents=[],
                error_code="PPTX_PARSE_FAILED",
                error_message=str(exc)[:500],
            )


register_parser(PptxParser())
